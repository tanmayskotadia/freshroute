"""Generate the faculty handbook as editable Word and Markdown from one content source."""
import json, re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from content import build

root=Path(__file__).resolve().parents[2]
out=root/'output/faculty-guide'; out.mkdir(parents=True,exist_ok=True)
notes=root/'docs/faculty-guide'; notes.mkdir(parents=True,exist_ok=True)
seed=json.loads((root/'docs/phase-3/seeded-incident.json').read_text(encoding='utf-8'))
pages=build(seed)
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=sec.footer_distance=Inches(.25)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style=doc.styles[name];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
    for color in style.element.xpath('.//w:color'):
        for attr in ['themeColor','themeTint','themeShade']:
            if color.get(qn('w:'+attr)) is not None:del color.attrib[qn('w:'+attr)]
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
for name,size in [('Title',28),('Heading 1',22),('Heading 2',12.5)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_after=Pt(7)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Heading 2'].paragraph_format.space_before=Pt(9)
for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
header=sec.header.paragraphs[0];header.text='FreshRoute website and faculty demonstration guide'
for run in header.runs:run.font.size=Pt(9);run.font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('23 October 2026  |  Page ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
for run in footer.runs:run.font.size=Pt(9)
doc.core_properties.title='FreshRoute website and faculty demonstration guide'
doc.core_properties.subject='User instructions and faculty explanation of the implemented cloud architecture'
doc.core_properties.author='FreshRoute project'

md=[]
for index,item in enumerate(pages):
    title=doc.add_paragraph(item['title'],'Title' if index==0 else 'Heading 1')
    title.paragraph_format.page_break_before=bool(index)
    md.extend(['# '+item['title'],''])
    for b in item['blocks']:
        if b['kind']=='h':doc.add_paragraph(b['text'],'Heading 2');md.extend(['## '+b['text'],''])
        elif b['kind'] in ['p','say']:
            text=('Say to faculty: '+b['text']) if b['kind']=='say' else b['text']
            para=doc.add_paragraph(text)
            if b['kind']=='say':
                para.paragraph_format.left_indent=Inches(.15)
                for run in para.runs:run.italic=True
            if index==len(pages)-1:
                for run in para.runs:run.font.size=Pt(9.5)
                para.paragraph_format.space_after=Pt(4)
            md.extend([('> '+b['text']) if b['kind']=='say' else b['text'],''])
        elif b['kind']=='steps':
            for i,txt in enumerate(b['items'],1):
                para=doc.add_paragraph(f'{i}. {txt}')
                para.paragraph_format.left_indent=Inches(.18)
                para.paragraph_format.first_line_indent=Inches(-.18)
                para.paragraph_format.space_after=Pt(5)
                md.append(f'{i}. {txt}')
            md.append('')
        elif b['kind']=='diagram':
            diagram=root/'output/phase-9/FreshRoute-Architecture.png'
            if diagram.exists():
                para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                para.add_run().add_picture(str(diagram),width=Inches(4.8))
                para.paragraph_format.space_after=Pt(2)
                cap=doc.add_paragraph('Simplified retained application flow from the editable project presentation.')
                for run in cap.runs:run.font.size=Pt(9)
            md.extend(['Architecture diagram: see the project presentation and docs/phase-9/ARCHITECTURE.md.',''])
        elif b['kind']=='table':
            rows=[b['headers'],*b['rows']]
            tab=doc.add_table(rows=0,cols=len(b['headers']));tab.autofit=False
            widths=b.get('widths') or [6.7/len(b['headers'])]*len(b['headers'])
            for col,width in zip(tab.columns,widths):col.width=Inches(width)
            props=tab._tbl.tblPr
            borders=OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge)
                for key,val in [('val','single'),('sz','5'),('color','D9D9D9')]:el.set(qn('w:'+key),val)
                borders.append(el)
            props.append(borders)
            for rnum,values in enumerate(rows):
                row=tab.add_row();rp=row._tr.get_or_add_trPr()
                no_split=OxmlElement('w:cantSplit');rp.append(no_split)
                if rnum==0:
                    repeat=OxmlElement('w:tblHeader');rp.append(repeat)
                for j,(cell,value) in enumerate(zip(row.cells,values)):
                    cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text=str(value);cp=cell._tc.get_or_add_tcPr()
                    margins=OxmlElement('w:tcMar')
                    for edge,value in [('top','65'),('bottom','65'),('left','100'),('right','100')]:
                        el=OxmlElement('w:'+edge);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');margins.append(el)
                    cp.append(margins)
                    if rnum==0 or rnum%2==0:
                        shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'303030' if rnum==0 else 'F5F5F5');cp.append(shade)
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1)
                        para.paragraph_format.line_spacing=1.06
                        for run in para.runs:
                            run.font.size=Pt(10);run.bold=rnum==0;run.font.color.rgb=RGBColor(255,255,255) if rnum==0 else RGBColor(0,0,0)
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
            md.extend(['| '+' | '.join(b['headers'])+' |','| '+' | '.join(['---']*len(b['headers']))+' |'])
            md.extend('| '+' | '.join(map(str,r))+' |' for r in b['rows']);md.append('')
doc.save(out/'FreshRoute-Website-and-Faculty-Guide.docx')
(notes/'GUIDE.md').write_text('\n'.join(md),encoding='utf-8')
(notes/'CONTENT.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')
stats={'plannedPages':len(pages),'sections':len(pages)-2,'words':len(re.findall(r'\S+',' '.join(md))),'output':str(out/'FreshRoute-Website-and-Faculty-Guide.docx')}
(notes/'BUILD.json').write_text(json.dumps(stats,indent=2),encoding='utf-8')
print(json.dumps(stats))
