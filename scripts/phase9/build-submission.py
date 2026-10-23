"""Rebuild the evaluation documents from project evidence; contains no credentials."""
import csv, json, html, hashlib
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'output/phase-9'; DOCS=ROOT/'docs/phase-9'
OUT.mkdir(parents=True,exist_ok=True); DOCS.mkdir(parents=True,exist_ok=True)
seed=json.loads((ROOT/'docs/phase-3/seeded-incident.json').read_text(encoding='utf-8'))
candidates=seed['result']['candidates']; winner=next(c for c in candidates if c['id']==seed['result']['winnerId'])
baseline=next(c for c in candidates if c['id']=='continue')

# Each named syllabus item has a row. Status describes the actual project boundary.
modules={
1: [('Introduction to Cloud Computing','Core','On-demand HTTP API, Lambda and object storage'),('Overview of AWS Global Infrastructure','Study','Mumbai region, AZ and edge scope; no multi-region test'),('AWS Shared Responsibility Model','Core + study','Customer IAM/data controls versus AWS infrastructure'),('AWS Identity and Access Management (IAM)','Core','Scoped execution roles'),('IAM Users','Study','Human access design; no unnecessary long-lived user keys'),('IAM Groups','Study','Human group design, separate from Cognito app identity'),('IAM Policies','Core','Resource and owner-prefix restrictions'),('IAM Roles','Core','Lambda, workflow and service trust policies'),('IAM Best Practices','Core + study','Least privilege and short-lived identity; human MFA discussed')],
2: [('AWS Compute Services','Core + lab deleted','Lambda product and bounded EC2 experiment'),('Amazon EC2','Lab deleted','Two workers and SSM command evidence'),('EC2 Instance Types','Lab + study','t3.micro chosen; workload sizing comparison'),('EC2 Pricing Models','Study','On-Demand, Spot and commitments; no commitments purchased'),('EC2 Instance Lifecycle','Lab deleted','Launch, terminate and replacement observed; no stop/start claim'),('Auto Scaling','Lab deleted','Controlled 1-to-2 scale-out and replacement'),('Elastic Load Balancing','Lab deleted','Internal ALB and healthy target checks'),('AWS Storage Services','Core + lab deleted','Object, block and shared-file boundaries'),('Amazon S3','Core','Four private versioned buckets'),('S3 Storage Classes','Study','Hot Standard objects; archive trade-offs'),('S3 Lifecycle Policies','Core','Configured expiry/noncurrent/multipart policies; not archive proof'),('Amazon EBS','Lab deleted','Distinct per-worker scratch content'),('Amazon Glacier','Study','Archive and restore design; no restore executed'),('Amazon EFS','Lab deleted','Same shared-file checksum on both workers')],
3: [('AWS Networking and Content Delivery','Core + lab deleted','HTTPS fallback plus isolated network experiments'),('Amazon VPC','Lab deleted','Two non-overlapping lab VPCs'),('VPC Subnets','Lab deleted','Two-AZ compute lab and private SQL subnets'),('Route Tables','Lab deleted','Scoped return routes for peering'),('Security Groups','Lab deleted','Allowed connectivity and denied SSH'),('VPC Peering','Lab deleted','Bidirectional connectivity checks'),('AWS Direct Connect','Study','Warehouse hybrid-link design; no physical circuit'),('Amazon CloudFront','Blocked / study','Account verification prevents creation; approved API Gateway HTTPS fallback')],
4: [('AWS Database Services','Core + lab deleted','DynamoDB live; relational fixture lab'),('Amazon RDS','Lab deleted','Private PostgreSQL fixture and verified SQL TLS'),('RDS Multi-AZ Deployments','Study','Standby/failover design; no DB failover test'),('RDS Read Replicas','Study','Read routing and lag comparison'),('Amazon Aurora','Study','Relational cluster alternative'),('Amazon DynamoDB','Core','Four tables, conditional writes and capacity transactions'),('DynamoDB Streams','Core','Decision outbox event to background worker')],
5: [('AWS Management Tools','Core + lab deleted','CDK/CloudFormation, monitoring and managed operations'),('AWS CloudWatch','Core','Dashboard, 14 alarms and canary state transitions'),('AWS CloudTrail','Core','Encrypted audit objects/digests and selected actual API event'),('AWS Config','Lab deleted','Compliant, drift, remediation and recheck'),('AWS Systems Manager','Lab deleted','Managed Run Command without SSH'),('AWS Cost Management','Core','Dated estimated Cost Explorer snapshot and cleanup evidence'),('AWS Free Tier','Account review','Dated account credits and expiry; not all services free'),('AWS Billing and Account Management','Account review','Resource ownership and delayed gross-cost view'),('AWS Budgets','Core','Five gross-cost thresholds; not a spending cap'),('AWS Trusted Advisor','Restricted / study','Support API subscription required; no upgrade purchased')],
6: [('AWS Security and Compliance','Core + study','Controls demonstrated; no compliance certification'),('AWS Shared Responsibility Model','Core + study','Service-specific customer control allocation'),('AWS Key Management Service (KMS)','Core','Rotating customer-managed key for reports and audit'),('AWS WAF','Lab deleted','Benign test 403 versus allowed 200 on lab ALB'),('AWS Shield','Study','DDoS protection comparison; no Advanced subscription/test'),('AWS Security Best Practices','Core','TLS, owner checks, private S3, input validation, bounded AI')],
7: [('AWS Application Integration','Core','Event-driven import, analysis and post-decision jobs'),('Amazon SQS','Core','Import/post-decision queues, DLQs and failure queue'),('Amazon SNS','Core; email unverified','Topics and publish evidence; rescue email confirmation still unverified'),('Amazon SWF','Study','Worker coordination compared with Step Functions'),('AWS Step Functions','Core','Standard analysis workflow with retry/catch'),('AWS Well-Architected Framework','Review','Six-pillar project review and trade-offs'),('Design Principles for AWS Cloud Architectures','Review','Managed services, loose coupling and explicit failure paths'),('Scalability and Elasticity','Core + lab deleted','Queue/concurrency design and controlled ASG scale-out'),('High Availability and Fault Tolerance','Core + lab deleted','Managed-service design plus observed instance replacement; no uptime SLA'),('Security','Core','Owner-bound APIs and private artifacts'),('Cost Optimization','Core','Deleted labs, bounded tokens, quotas and retention')],
8: [('AWS Migration and Hybrid Architecture','Study + lab deleted','Warehouse migration design and isolated reconciliation'),('AWS Database Migration Service (DMS)','Lab deleted','Full load, 100%, one table, zero errors, three exact rows'),('AWS Server Migration Service (SMS)','Historical study','Syllabus comparison only; not provisioned'),('AWS Snowball','Study','Offline transfer scenario; not ordered'),('AWS Snowball Edge','Study','Edge processing/transfer scenario; not ordered'),('AWS Snowmobile','Historical study','Large-scale offline transfer comparison; not provisioned')]
}
rows=[(m,*item) for m,items in modules.items() for item in items]
with (DOCS/'SYLLABUS_MATRIX.csv').open('w',newline='',encoding='utf-8') as f:
    w=csv.writer(f); w.writerow(['Module','Exact syllabus topic','Actual treatment','Evidence or explanation']); w.writerows(rows)

def page(title,*paras,table=None): return {'title':title,'paras':list(paras),'table':table}
pages=[
page('FreshRoute', 'Cloud Architecture Design | BCSE355L | Phase 9 submission | 23 October 2026',
     'A cloud decision-support platform for perishable shipments. Observe a delay, compare practical rescue options, preserve the calculation, and review a capacity-backed operational decision.',
     'Submission contents: technical report, editable presentation, viva cards, seven-minute demo script, syllabus matrix, source and deployment guide, recorded AWS evidence, and a captioned architecture overview video.',
     'Deployment: AWS account 118462784511, Mumbai (ap-south-1). Website: https://dr4sw5bpoi.execute-api.ap-south-1.amazonaws.com',
     'Evidence boundary: the application and monitoring remain deployed. Disposable engineering labs were verified and deleted. Studies are labelled individually. This academic model is not a food-safety assessment.'),
page('1. Problem, users and practical value',
     'A delayed produce shipment creates a time-sensitive choice: continue, cool the shipment, or reroute to a receiving location. A dispatcher needs a reproducible comparison of likely quality, travel assumptions, cost, receiving hours and available capacity.',
     'FreshRoute serves a fictional distributor operator and an administrator who manages receiving rules. Operators inspect shipments, calculate alternatives, save analyses and explicitly approve operational decisions. Administrators control destination rules. The cloud API checks ownership on every relevant object.',
     'The distinctive feature is the connection between explanation and execution. A saved calculation is immutable; a later reservation checks current versions and capacity. AI explains calculated facts and proposes drafts, while the application owns the arithmetic and transaction.',
     'The intended real-life use is dispatcher decision support. There is no live GPS, IoT sensor feed, traffic prediction, recipient marketplace, logistics dispatch integration or verified food-safety guarantee. Shipment observations and route assumptions are entered by the user.'),
page('2. Product journey and interface',
     'The forest-and-cream interface groups the journey into Observe, Simulate and Decide. Overview and Shipments expose search, filters and incident details. Shipment forms preserve drafts; receiving rules and entered temperature segments remain editable with validation.',
     'Rescue Lab compares continuation, cooling and alternative destinations, explains rejected options and shows sensitivity to changed assumptions. History reopens saved inputs/results. Split planning conserves quantity and checks aggregate budget and destination capacity.',
     'The AWS workspace adds Cognito sign-in, cloud persistence, CSV imports, background analysis jobs, AI explanations and scenario drafts, reservations, private PDF reports, notification records and outcome analytics. Local mode supports browser-local calculation workflows without AWS credentials.',
     'Phase 9 adds a collapsible seven-minute presentation guide with screen navigation, segment timings and a pauseable clock. It never automatically creates a reservation. Keyboard focus filtering now skips disabled/hidden dialog controls and includes links.',
     'Visual QA limitation: the existing browser-origin policy blocked live preview checks. Production compilation and source review passed, but phone/desktop screenshots and a human keyboard rehearsal remain pre-submission checks. Earlier screenshots are dated Phase 3 evidence, not current captures.'),
page('3. Retained cloud architecture',
     'Browser requests reach an HTTPS API Gateway frontend and the Web Lambda, which reads the private web bucket. CloudFront creation is blocked by account verification; the user approved this HTTPS fallback. API Gateway HTTP API routes authenticated requests to application handlers using Cognito access tokens.',
     'Eight Lambda functions separate Web, Workspace API, Verification, Import Worker, Analysis Worker, Outbox Bridge, Post-decision Worker and Copilot responsibilities. Four DynamoDB tables store Shipments, Runs, Destinations and Activity with owner-aware keys, on-demand billing and point-in-time recovery.',
     'Step Functions Standard coordinates analysis. SQS separates import and post-decision work with dead-letter handling. DynamoDB Streams drives an outbox bridge. SNS supports notifications; CloudWatch observes operations. CloudTrail delivers audit objects and digests to encrypted storage.',
     'The application remains useful after lab deletion. EC2, ALB, EFS, RDS, DMS, Config and WAF experiments are evidence of syllabus learning rather than dependencies of the retained dispatcher application.'),
page('4. Storage and data boundaries',
     'Four S3 buckets exist for different readers, writers and retention needs. The web bucket stores release assets; data stores uploads and fixtures; reports stores generated owner-bound downloads; audit receives CloudTrail evidence. All are private, versioned and require TLS. A customer-managed KMS key protects reports and audit data.',
     'DynamoDB owner partition keys and object identifiers support per-user access. Activity stores jobs, notification records and AI cache/quota items; the JobsByOwnerTime index supports history queries. TTL is used for expiring cache-related records and is not an immediate delete guarantee.',
     'Capacity uses whole grams; money uses integer paise. Conditional transactions prevent a stale calculation or duplicate submission from silently consuming capacity twice. A terminal transition releases capacity once under the allowed state machine.',
     'Lifecycle rules are configuration evidence. Tiny hot objects remain in Standard; no actual Glacier transition or restore is claimed. Four buckets are justified by access/retention boundaries, not by multiplying storage for presentation value.'),
page('5. Deterministic model and worked example',
     'The illustrative Q10 model converts each entered temperature segment into equivalent hours at 10 C: effectiveHours = sum(hours * 2^((temperatureC - 10) / 10)). Each quality cohort then decays exponentially using produce-specific teaching parameters.',
     'Estimated loss sums the mass of cohorts below the declared quality threshold; it is not quantity multiplied by one minus average quality. Feasibility checks receiving hours, accepted produce, capacity, budget and modeled quality before weighted ranking. Route times and costs are declared assumptions.',
     f"The unchanged FR-101 seed contains 500 kg of tomatoes. Continuation has modeled quality {baseline['qualityIndex']:.1f}/100 and {baseline['estimatedLossKg']:.0f} kg below threshold; it is infeasible. The seeded winner is {winner['destinationName']} with modeled quality {winner['qualityIndex']:.1f}/100, estimated loss {winner['estimatedLossKg']:.0f} kg and declared transport cost INR {winner['costPaise']/100:.0f}.",
     'These are reproducible outputs of docs/phase-3/seeded-incident.json, not measured waste savings or a commercial forecast. A saved snapshot preserves model version, source versions, input signature and results. Later shipment/rule edits require a new calculation before confirmation.'),
page('6. AI integration and operator control',
     'The neutral AI interface explains a saved analysis and proposes a structured scenario draft. The current backend uses the Groq API, model openai/gpt-oss-20b, through a server-side adapter. A Bedrock adapter remains optional; successful Bedrock use is not claimed.',
     'The API credential stays in Secrets Manager and resolves through a CloudFormation dynamic reference at deployment. It is never included in the browser bundle or submission. The runtime request has an 800-token limit and a nine-second deadline.',
     'Structured output can reference only authoritative fact identifiers and validated draft fields. Unsupported output, timeout or missing availability falls back to deterministic explanations. Cache entries last one hour and include owner, model/provider, question, run and prompt version in the key.',
     'Daily uncached request limits are 20 per owner and 200 globally in UTC. Failed uncached attempts count. IAM limits AI writes to AI-prefixed state, so it cannot create decisions or reserve capacity. Applying a draft and confirming a reservation are separate explicit operator actions.'),
page('7. Operational reliability and security',
     'Cognito PKCE avoids a browser client secret. The API verifies the access token and owner context. Private reports are downloaded through short-lived links after authorization; direct anonymous S3 access is denied. Backend validation recomputes authoritative results rather than trusting a browser-supplied score.',
     'Imports and post-decision workloads have independent queues and dead-letter handling. Retries use idempotent business transitions; duplicate events must not double-book or release capacity twice. The isolated reliability lab observed poison-message redrive and corrected replay without consuming production DLQ messages.',
     'CloudWatch has a dashboard and 14 alarms. A canary moved OK to ALARM and back to OK with successful SNS actions. CloudTrail delivered encrypted objects and digests, and an actual Auto Scaling termination API event was selected. Digest-chain cryptographic validation was not executed.',
     'Email delivery remains a limitation: the earlier rescue subscription was still pending at its last verification, and the operational topic has no email subscriber. In-app notification records and SNS publish success are distinct from confirmed human email delivery. No compliance certification or uptime SLA is asserted.'),
page('8. Engineering labs: evidence and limits',
     'Compute/network: two healthy workers behind an internal ALB; equal EFS file checksums and distinct EBS scratch content; bidirectional VPC peering; denied SSH; Systems Manager Run Command. Terminating a worker produced a new healthy replacement with the same shared-file checksum. Controlled scaling changed desired capacity from one to two.',
     'Security/configuration: WAF returned 403 for a benign test pattern while the ordinary request returned 200. A temporary security-group drift moved the Config check from COMPLIANT to NON_COMPLIANT, then back after remediation. These tests do not demonstrate WAF on the main HTTP API or DDoS resistance.',
     'SQL/migration: private PostgreSQL accepted a client connection with verified TLS 1.3. DMS full load reached 100%, one table loaded and zero tables errored; three source and target rows matched exactly. Source and target were schemas on the same disposable RDS instance. There was no on-premises migration, CDC or production cutover.',
     'All five stacks reached DELETE_COMPLETE: ComputeLab, SqlLab, MigrationLab, ConfigLab and ReliabilityLab. Inventory checked absent lab compute, storage, network, RDS, DMS, queues and recorder. Retained Foundation, Application and Operations stacks remain deployed.'),
page('9. Syllabus coverage: modules 1-4',
     'The full CSV contains one row for every named topic in the supplied BCSE355L syllabus. Coverage is not a claim that every service is active.',
     table=[['Module','Implemented evidence','Studies / limits'],['1: Foundations & IAM','Execution roles, policies, ownership and shared controls','Global infrastructure; human users/groups and MFA design'],['2: Compute & storage','Lambda; four S3 buckets; deleted EC2/ASG/ALB/EBS/EFS lab','Instance/pricing comparison; storage classes/Glacier; no archive restore'],['3: Networking','VPCs, subnets, routes, SGs and peering lab; HTTPS product','Direct Connect study; CloudFront blocked by verification'],['4: Databases','DynamoDB/Streams product; private RDS fixture lab','Multi-AZ, read replicas and Aurora studies; no DB failover test']]),
page('10. Syllabus coverage: modules 5-8',
     'Service coverage is supported by retained resources, dated experimental evidence or a labelled design comparison. Historical syllabus entries are discussed without implying provisioning or current availability.',
     table=[['Module','Implemented evidence','Studies / limits'],['5: Management & cost','CloudWatch, CloudTrail, Budgets; deleted Config/SSM lab; account review','Trusted Advisor support API restricted; delayed billing'],['6: Security','KMS, TLS, owner checks; WAF lab and Config remediation','Shield comparison; no Advanced subscription or compliance certification'],['7: Integration & design','SQS/SNS, Streams, Step Functions; replacement/scale tests','SWF study; Well-Architected review; no human email delivery claim'],['8: Migration & hybrid','DMS full load with exact three-row reconciliation','SMS, Snowball, Snowball Edge, Snowmobile and hybrid design studies']]),
page('11. Cost, lifecycle and Well-Architected review',
     'USD 200 is the project planning ceiling, not a hard cap. Budgets thresholds are USD 25, 50, 100, 150 and 180; notification latency means the user must still monitor and delete resources. Lab deletion limits ongoing spend, while retained application, logs, key and storage continue normal charges.',
     'Dated account evidence showed USD 179.94 remaining credits and plan expiry 28 February 2027. A 22 October UTC account-wide estimated gross Cost Explorer snapshot excluding credits/refunds reported USD 0.0000150018 for 1-22 October. Billing is delayed and the figure is not isolated project spend or a final invoice.',
     'Operational excellence: versioned infrastructure, evidence and runbooks. Security: owner checks, private storage and scoped roles. Reliability: idempotence, queues, catch paths and observed replacement. Performance efficiency: managed on-demand services and bounded AI requests.',
     'Cost optimization: deleted labs, quotas and lifecycle rules. Sustainability: retain the useful application and remove idle teaching infrastructure. Production expansion would require calibrated quality models, stronger operational recovery evidence, recipient integration, load tests and confirmed notification delivery.'),
page('12. Reproduce and demonstrate',
     'Local setup: install Node 24, run npm.cmd ci, npm.cmd test and npm.cmd run build, then npm.cmd run dev. Open the displayed localhost URL. No AWS credentials or AI secret is needed for the local demo. The ZIP excludes node_modules, secrets, signed URLs, temporary transfer files and the .git directory.',
     'Cloud setup is documented in RUN_DEPLOY_GUIDE.md and earlier phase deployment notes. CDK synthesis describes Foundation/Application; Operations and lab templates have their own scripts. Review change sets, retain-state policies and the approved hosting fallback before deployment. Do not recreate deleted labs merely to open the website.',
     'Demonstration: incident (45 s), calculation (85 s), frozen history (60 s), AI (60 s), fulfilment (70 s), architecture/evidence (70 s), value/limitations (30 s). Total 420 seconds. Presentation mode provides notes and navigation; a human should rehearse with the saved demo data before evaluation.',
     'Fresh Phase 9 checks cover production compilation, 49 unit tests, server type checks, isolated install/build verification and cloud hosting smoke checks. The earlier 19 live application checks remain dated Phase 8 regression evidence; they must not be relabelled as a new Phase 9 run.'),
page('13. Evidence register and references',
     'Primary requirements: supplied CLOUD-ARCHITECTURE-DESIGN.pdf, BCSE355L syllabus version 2.0, modules 1-8. The submission matrix was checked against the supplied text, including historical migration/transfer topics.',
     'Implementation: src/model/engine.ts and parameters.ts; server/ handlers and probes; infrastructure/ CDK; scripts/ operations and laboratory templates. Phase 3 seeded-incident.json supports the worked example. docs/phase-7/ documents AI behavior and actual provider configuration.',
     'Measured cloud evidence: docs/phase-8/application-regression.json; compute-verification.json; replacement-verification.json; scale-and-migration-completion.json; config-verification.json; reliability-verification.json; migration-verification.json; operations-final-verification.json; cleanup-verification.json; cost-verification.json and account-review.json.',
     'Phase 9 evidence and reproduction: docs/phase-9/VERIFICATION.json, EVIDENCE_INDEX.md, RUN_DEPLOY_GUIDE.md, DEMO_SCRIPT.md and SYLLABUS_MATRIX.csv. Machine evidence records actual observations; a successful API call alone is not proof of all service behavior.',
     'The report intentionally records unresolved limits: browser visual verification, CloudFront account verification, email subscription/delivery and human rehearsal. No grade, actual food rescued, commercial savings or food-safety outcome is guaranteed.')
]

styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='FRTitle',fontName='Helvetica-Bold',fontSize=25,leading=30,textColor=colors.black,spaceAfter=22))
styles.add(ParagraphStyle(name='FRBody',fontName='Helvetica',fontSize=11,leading=17,spaceAfter=15))
styles.add(ParagraphStyle(name='FRCell',fontName='Helvetica',fontSize=9,leading=13))
def para(s,style='FRBody'): return Paragraph(html.escape(s),styles[style])
def footer(canvas,doc):
    canvas.setStrokeColor(colors.HexColor('#ced6cd')); canvas.line(54,43,558,43)
    canvas.setFillColor(colors.HexColor('#536354')); canvas.setFont('Helvetica',8)
    canvas.drawString(54,30,'FreshRoute | Cloud Architecture Design | 23 October 2026')
    canvas.drawRightString(558,30,str(doc.page))

doc=Document(); section=doc.sections[0]; section.page_width=Inches(8.5); section.page_height=Inches(11)
section.top_margin=section.bottom_margin=Inches(.7); section.left_margin=section.right_margin=Inches(.75)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11); doc.styles['Normal'].paragraph_format.space_after=Pt(12)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Title'].font.size=Pt(28); doc.styles['Heading 1'].font.size=Pt(22)
foot=section.footer.paragraphs[0]; foot.text='FreshRoute | BCSE355L | Phase 9     '
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
story=[]
for i,p in enumerate(pages):
    if i: doc.add_page_break(); story.append(PageBreak())
    doc.add_paragraph(p['title'],'Title' if i==0 else 'Heading 1'); story.append(para(p['title'],'FRTitle'))
    for s in p['paras']: doc.add_paragraph(s); story.append(para(s))
    if p['table']:
        tab=doc.add_table(rows=0,cols=3); tab.style='Table Grid'
        for ri,row in enumerate(p['table']):
            cells=tab.add_row().cells
            for c,v in zip(cells,row):
                c.text=v
                for run in c.paragraphs[0].runs: run.font.size=Pt(10); run.bold=ri==0
        t=Table([[para(x,'FRCell') for x in row] for row in p['table']],colWidths=[105,198,201],hAlign='LEFT')
        t.setStyle(TableStyle([('GRID',(0,0),(-1,-1),.5,colors.HexColor('#b9c5b9')),('BACKGROUND',(0,0),(-1,0),colors.HexColor('#edf1eb')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),9),('RIGHTPADDING',(0,0),(-1,-1),9),('TOPPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,0),(-1,-1),10)]));story.append(t)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.save(OUT/'FreshRoute-Technical-Report.docx')
SimpleDocTemplate(str(OUT/'FreshRoute-Technical-Report.pdf'),pagesize=(612,792),leftMargin=54,rightMargin=54,topMargin=54,bottomMargin=60).build(story,onFirstPage=footer,onLaterPages=footer)
(DOCS/'REPORT_SOURCE.json').write_text(json.dumps(pages,indent=2),encoding='utf-8')

qa=[
('What is FreshRoute?','A dispatcher decision-support app that compares rescue options for delayed perishable shipments, saves the calculation, and checks capacity before an operator confirms a plan.'),
('Why is it useful?','It connects an explainable comparison with an auditable operational decision. It is a prototype: no measured food savings or recipient integration is claimed.'),
('Why four S3 buckets?','Web, data, reports and audit have different writers, readers and retention needs. All are private and versioned; reports and audit use the customer-managed KMS key.'),
('Why DynamoDB instead of RDS for the main app?','The access patterns are owner-and-ID lookups, histories and conditional capacity transactions. RDS was an isolated legacy SQL teaching lab, not a second live source of truth.'),
('How do you prevent double-booking?','Whole-gram quantities, conditional versions, transaction checks, stable operation identity and allowed state transitions protect capacity from duplicate or stale requests.'),
('Can AI reserve a shipment?','No. It returns validated explanations or drafts. Its IAM writes are limited to AI state; the operator must explicitly review and confirm business actions.'),
('Which model provider is used?','The current backend adapter uses Groq openai/gpt-oss-20b. The interface uses neutral AI wording. Bedrock is an optional adapter and successful Bedrock use is not claimed.'),
('Where is the API key?','Secrets Manager, freshroute/copilot-api-key, field apiKey. CloudFormation resolves it for the server runtime. Never expose it in frontend config, source, slides or logs.'),
('What happens when AI is unavailable?','A nine-second deadline, strict schema/fact validation and deterministic fallback keep the saved calculation authoritative. Failure does not approve or reserve anything.'),
('How is AI cost bounded?','800 output tokens, a one-hour owner-aware cache and UTC uncached request quotas of 20 per owner and 200 globally per day. Failed uncached attempts count.'),
('How is quality calculated?','Entered temperatures become equivalent exposure at 10 C using Q10=2; produce-specific exponential decay affects weighted quality cohorts. Parameters are illustrative.'),
('How is estimated loss calculated?','Sum cohort mass whose modeled quality is below the declared threshold. It is not total mass times one minus mean quality, and it is not measured spoilage.'),
('Does the top score always win?','Hard constraints run before ranking. Closed receiving hours, rejected produce, low quality, excess budget or insufficient capacity can make a high score infeasible.'),
('Why freeze analyses?','A saved snapshot preserves inputs, model version and results for comparison. Later source edits cannot rewrite history and require recalculation before confirmation.'),
('Why use Step Functions?','The Standard workflow exposes analysis steps, retries, catch paths and job state instead of hiding coordination inside one long browser request.'),
('Why separate queues?','Imports and post-decision jobs have different work and failure patterns. Each can retry or fail without tying up the other workload.'),
('What is a DLQ?','A dead-letter queue preserves repeatedly failing messages for diagnosis. The isolated poison/replay test used only lab queues; production DLQ contents were preserved.'),
('What does DynamoDB Streams do?','A decision-state change creates a downstream event for the outbox bridge. Consumers must handle duplicate delivery idempotently.'),
('SNS versus SQS?','SNS publishes events to subscribers; SQS buffers work for consumers. Publish success is not evidence that a person received an email.'),
('Has email delivery been proved?','No new human-delivery claim is made. Rescue subscription confirmation remained unverified at the last check; the operations topic has no email subscriber.'),
('Why no CloudFront?','Account verification blocked distribution creation. The approved fallback is HTTPS API Gateway plus a Web Lambda reading private S3. CloudFront stays labelled blocked/study.'),
('Is the EC2 lab part of the production path?','No. Two workers, internal ALB, EFS/EBS and peering were bounded teaching experiments. They were deleted after evidence capture.'),
('EBS versus EFS versus S3?','EBS is block storage attached to compute; EFS is a shared filesystem; S3 stores objects through an API. The lab observed distinct scratch bytes and equal shared-file checksums.'),
('What proves Auto Scaling?','The evidence records a terminated worker replaced by a healthy new worker and controlled scale-out from desired one to two. It does not prove arbitrary traffic-load scalability.'),
('What proves networking?','Allowed peering connectivity with return routes, denied SSH, managed SSM commands, and private SQL access. VPC peering is not a physical Direct Connect link.'),
('Multi-AZ versus read replica?','A standby targets availability/failover; a read replica targets read scaling with replication lag. Both are studies here; the SQL lab was Single-AZ.'),
('What exactly did DMS migrate?','Three fictional rows between two schemas on the same disposable PostgreSQL RDS instance. Full load completed at 100% with zero table errors and exact row reconciliation.'),
('Did you test CDC or production cutover?','No. Full load and reconciliation were tested. CDC, external-source migration, cutover and rollback remain designs, not measured outcomes.'),
('What did WAF prove?','An ordinary lab request returned 200 and a benign known pattern returned 403. The WAF protected the compatible lab ALB, not the main HTTP API.'),
('CloudWatch, CloudTrail and Config?','CloudWatch observes logs/metrics/alarms; CloudTrail records API activity; Config tracks configuration/compliance. Separate evidence supports each role.'),
('What is the KMS evidence?','Reports and audit objects use the customer-managed key with rotation and scoped permissions. A delivered encrypted audit object was checked; no certification is implied.'),
('Did you validate CloudTrail digest chains?','Delivery of audit objects and digests was observed, and validation was enabled. Cryptographic digest-chain verification was not executed.'),
('What did Config remediation show?','A narrow temporary security-group drift changed compliance, then remediation restored COMPLIANT. The temporary recorder and lab resources were deleted.'),
('Is everything free?','No. Free Tier and credits are account-specific and time-bound. USD 200 is a planning ceiling; budgets alert after delayed usage rather than enforcing a hard stop.'),
('What remains chargeable?','Retained application, storage, encryption key, logs, monitoring and requests can accrue charges. Deleted lab inventory is documented; billing estimates are delayed.'),
('How are all syllabus services covered?','The CSV maps every named topic to retained implementation, verified/deleted lab, restricted service or study. Snow family, SMS, SWF, Aurora and Direct Connect are not shown as active.'),
('What are the six Well-Architected pillars here?','Operational excellence: runbooks; security: owner checks; reliability: retries/idempotence; performance: managed capacity; cost: lab cleanup; sustainability: avoid idle infrastructure.'),
('What does local mode prove?','Browser-local forms, persistence, deterministic calculation and snapshot workflows. It does not prove live authentication, cloud reservations, AI invocation or email.'),
('What would production need?','Calibrated models, real recipient/dispatch integration, confirmed delivery channels, recovery and load tests, monitoring ownership, data protection review and operational runbooks.'),
('What remains before presenting?','Human rehearsal and browser visual/keyboard review, account verification for CloudFront if desired, and subscription confirmation plus delivery testing before claiming email receipt.')]
vstory=[]
for i in range(0,len(qa),5):
    if i:vstory.append(PageBreak())
    vstory.append(para('Viva cards | '+str(i//5+1),'FRTitle'))
    for j,(q,a) in enumerate(qa[i:i+5],i+1):
        vstory.append(Paragraph(f'<b>{j:02}. {html.escape(q)}</b>',styles['FRBody']))
        vstory.append(para(a))
SimpleDocTemplate(str(OUT/'FreshRoute-Viva-Cards.pdf'),pagesize=(612,792),leftMargin=54,rightMargin=54,topMargin=45,bottomMargin=55).build(vstory,onFirstPage=footer,onLaterPages=footer)
(DOCS/'VIVA_CARDS.md').write_text('# FreshRoute viva cards\n\n'+'\n\n'.join(f'### {i}. {q}\n\n{a}' for i,(q,a) in enumerate(qa,1)),encoding='utf-8')
print(json.dumps({'reportSections':len(pages),'syllabusRows':len(rows),'vivaQuestions':len(qa),'seedWinner':winner['destinationName']}))

